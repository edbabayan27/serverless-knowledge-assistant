"""Compare chunking strategies and retrieval methods on the evaluation questions.

Every chunking strategy is combined with every retrieval method (keyword search, vector search,
and the two hybrid methods). For each question the top results are retrieved, and the question
counts as answered at rank N when the N-th retrieved text contains one of its evidence quotes.

The results are written to an Excel workbook. The per-question ranks are the raw data; every
recall, MRR and average in the summary sheets is an Excel formula over them.

Usage (from the repository root, after downloading the source PDF to .local/):
    uv run --group eval python scripts/eval_retrieval.py                    # local embedding model
    uv run --group eval python scripts/eval_retrieval.py --embedder titan   # Amazon Titan (Bedrock)
    uv run --group eval python scripts/eval_retrieval.py --embedder none    # keyword search only
"""

from __future__ import annotations

import argparse
import json
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from statistics import mean

import numpy as np
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from knowledge_assistant.chunking import (
    SECTION_SEPARATOR,
    Section,
    extract_sections,
    read_pdf,
    sentence_windows,
)
from knowledge_assistant.hybrid_search import Bm25Index, reciprocal_rank_fusion, weighted_fusion

SOURCE_PDF = Path(".local/wellarchitected-serverless-applications-lens.pdf")
SOURCE_URL = (
    "https://docs.aws.amazon.com/pdfs/wellarchitected/latest/serverless-applications-lens/"
    "wellarchitected-serverless-applications-lens.pdf"
)
QUESTIONS = Path("eval/questions.json")
RESULTS_DIR = Path("eval/results")

CANDIDATES = 30  # results taken from each search before fusion
MAX_RANK = 10  # a question not answered in the top 10 counts as missed
CONTEXT_K = 5  # number of retrieved texts that would be sent to the model
ALPHAS = (0.25, 0.5, 0.75)
RRF_K = 60
LOCAL_MODEL = "BAAI/bge-small-en-v1.5"
LOCAL_QUERY_PREFIX = "Represent this sentence for searching relevant passages: "


# --- Chunking strategies ---------------------------------------------------------------------


@dataclass
class Strategy:
    """Search items plus the context each item stands for (itself, or its parent section)."""

    name: str
    description: str
    size: int
    overlap: int
    headings_in_search_text: bool
    search_texts: dict[str, str] = field(default_factory=dict)  # item key -> text that is searched
    context_of: dict[str, str] = field(default_factory=dict)  # item key -> context id
    contexts: dict[str, str] = field(default_factory=dict)  # context id -> text sent to the model
    chunk_words: dict[str, int] = field(default_factory=dict)  # item key -> words in the chunk

    def add(self, key: str, text: str, heading: str, context_id: str, context: str) -> None:
        searchable = f"{heading}\n\n{text}" if self.headings_in_search_text and heading else text
        self.search_texts[key] = searchable
        self.chunk_words[key] = len(text.split())
        self.context_of[key] = context_id
        self.contexts[context_id] = context


def word_windows(length: int, size: int, overlap: int) -> list[tuple[int, int]]:
    """Fixed-size windows of words, ignoring sentences (the naive baseline)."""
    step = size - overlap
    return [
        (start, min(start + size, length)) for start in range(0, max(length - overlap, 1), step)
    ]


def build_strategies(sections: Sequence[Section]) -> list[Strategy]:
    def heading(section: Section) -> str:
        return SECTION_SEPARATOR.join(section.path)

    def chunked(name, description, size, overlap, windows, headings=True) -> Strategy:
        strategy = Strategy(name, description, size, overlap, headings)
        for s, section in enumerate(sections):
            for w, (start, end) in enumerate(windows(section.words, size, overlap)):
                text = " ".join(section.words[start:end])
                key = f"{s:03d}.{w:02d}"
                strategy.add(key, text, heading(section), key, text)
        return strategy

    every_word = [word for section in sections for word in section.words]
    fixed = Strategy(
        "A. Fixed size",
        "Windows of 250 words over the whole document; ignores sections and sentences",
        250,
        50,
        False,
    )
    for w, (start, end) in enumerate(word_windows(len(every_word), 250, 50)):
        text = " ".join(every_word[start:end])
        fixed.add(f"{w:03d}", text, "", f"{w:03d}", text)

    parent_child = Strategy(
        "F. Parent-child",
        "Search 120-word sentence chunks; send the whole section they belong to",
        120,
        25,
        True,
    )
    for s, section in enumerate(sections):
        section_text = " ".join(section.words)
        for w, (start, end) in enumerate(sentence_windows(section.words, 120, 25)):
            text = " ".join(section.words[start:end])
            parent_child.add(f"{s:03d}.{w:02d}", text, heading(section), f"{s:03d}", section_text)

    def words_only(words, size, overlap):
        return word_windows(len(words), size, overlap)

    return [
        fixed,
        chunked(
            "B. Sections + word windows",
            "Split by section, then 250-word windows (the first version)",
            250,
            50,
            words_only,
        ),
        chunked(
            "C. Sections + sentences",
            "Split by section, then whole sentences up to 250 words (current version)",
            250,
            50,
            sentence_windows,
        ),
        chunked(
            "D. C without headings",
            "Same chunks as C, but the heading path is not added to the searched text",
            250,
            50,
            sentence_windows,
            headings=False,
        ),
        chunked(
            "E. Small sentence chunks",
            "Split by section, then whole sentences up to 120 words",
            120,
            25,
            sentence_windows,
        ),
        parent_child,
    ]


# --- Embedding models ------------------------------------------------------------------------


class LocalEmbedder:
    """A small open embedding model run on this machine, as a stand-in for Titan."""

    name = f"Local model ({LOCAL_MODEL})"

    def __init__(self) -> None:
        from fastembed import TextEmbedding  # imported here so `--embedder none` doesn't need it

        self._model = TextEmbedding(LOCAL_MODEL)

    def documents(self, texts: Sequence[str]) -> np.ndarray:
        return _normalise(np.array(list(self._model.embed(list(texts)))))

    def query(self, text: str) -> np.ndarray:
        return _normalise(np.array(list(self._model.embed([LOCAL_QUERY_PREFIX + text]))))[0]


class TitanEmbedder:
    """Amazon Titan Text Embeddings V2 through Bedrock, the model the deployed pipeline uses."""

    name = "Amazon Titan Text Embeddings V2 (1024 dimensions)"

    def __init__(self, region: str) -> None:
        import boto3
        from botocore.config import Config

        from knowledge_assistant.config import DEFAULT_EMBEDDING_MODEL_ID
        from knowledge_assistant.embeddings import BedrockEmbedder

        client = boto3.client(
            "bedrock-runtime",
            region_name=region,
            config=Config(retries={"mode": "adaptive", "max_attempts": 10}),
        )
        self._embedder = BedrockEmbedder(
            client, model_id=DEFAULT_EMBEDDING_MODEL_ID, dimensions=1024
        )

    def documents(self, texts: Sequence[str]) -> np.ndarray:
        return _normalise(np.array([e.vector for e in self._embedder.embed_many(texts)]))

    def query(self, text: str) -> np.ndarray:
        return _normalise(np.array([self._embedder.embed(text).vector]))[0]


def _normalise(vectors: np.ndarray) -> np.ndarray:
    return vectors / np.linalg.norm(vectors, axis=1, keepdims=True)


# --- Running the evaluation ------------------------------------------------------------------

Ranking = list[tuple[str, float]]


def normalise_text(text: str) -> str:
    """Lowercase, single spaces and straight quotes, so evidence matching ignores formatting."""
    curly_quotes = {0x2018: "'", 0x2019: "'", 0x201C: '"', 0x201D: '"'}
    return " ".join(text.translate(curly_quotes).lower().split())


def evaluate(strategy: Strategy, questions: list[dict], embedder) -> list[dict]:
    keys = list(strategy.search_texts)
    bm25 = Bm25Index(strategy.search_texts)
    vectors = embedder.documents([strategy.search_texts[k] for k in keys]) if embedder else None
    contexts = {cid: normalise_text(text) for cid, text in strategy.contexts.items()}

    def vector_search(question: str) -> Ranking:
        scores = vectors @ embedder.query(question)
        best = np.argsort(-scores)[:CANDIDATES]
        return [(keys[i], float(scores[i])) for i in best]

    rows = []
    for q in questions:
        keyword = bm25.search(q["question"], top_k=CANDIDATES)
        rankings: dict[str, Ranking] = {"BM25 (keyword)": keyword}
        if embedder:
            vector = vector_search(q["question"])
            rankings["Vector"] = vector
            for alpha in ALPHAS:
                rankings[f"Hybrid weighted (alpha {alpha})"] = weighted_fusion(
                    vector, keyword, alpha=alpha
                )
            rankings[f"Hybrid RRF (k={RRF_K})"] = reciprocal_rank_fusion([vector, keyword], k=RRF_K)
        evidence = [normalise_text(e) for e in q["evidence"]]
        for method, ranking in rankings.items():
            # Several search items can belong to the same context (parent-child); keep the first.
            context_ids = list(dict.fromkeys(strategy.context_of[key] for key, _ in ranking))
            rank = next(
                (
                    position
                    for position, cid in enumerate(context_ids[:MAX_RANK], start=1)
                    if any(e in contexts[cid] for e in evidence)
                ),
                0,
            )
            top = context_ids[:CONTEXT_K]
            rows.append(
                {
                    "strategy": strategy.name,
                    "method": method,
                    "id": q["id"],
                    "type": q["type"],
                    "rank": rank,
                    "context_words": sum(len(strategy.contexts[cid].split()) for cid in top),
                }
            )
    return rows


# --- Excel workbook --------------------------------------------------------------------------

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
GREEN_SCALE = ColorScaleRule(
    start_type="num", start_value=0, start_color="F8696B",
    mid_type="num", mid_value=0.5, mid_color="FFEB84",
    end_type="num", end_value=1, end_color="63BE7B",
)  # fmt: skip


def write_table(sheet, top: int, headers: Sequence[str], rows: Sequence[Sequence]) -> None:
    for column, title in enumerate(headers, start=1):
        cell = sheet.cell(row=top, column=column, value=title)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, values in enumerate(rows, start=top + 1):
        for c, value in enumerate(values, start=1):
            sheet.cell(row=r, column=c, value=value).font = Font(name=FONT)


def set_widths(sheet, widths: Sequence[float]) -> None:
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width


def title(sheet, text: str, note: str) -> None:
    sheet["A1"] = text
    sheet["A1"].font = Font(name=FONT, bold=True, size=14)
    sheet["A2"] = note
    sheet["A2"].font = Font(name=FONT, italic=True, color="595959")


def build_workbook(strategies, methods, questions, rows, embedder_name: str) -> Workbook:
    workbook = Workbook()
    n = len(rows)
    last = n + 1  # data rows on 'Per question' are 2..last

    # Per question: the raw data, plus hit flags computed from the rank.
    detail = workbook.active
    detail.title = "Per question"
    write_table(
        detail,
        1,
        ["Strategy", "Retrieval", "Question", "Type", "Rank (0 = not in top 10)", "Hit@1",
         "Hit@3", "Hit@5", "Reciprocal rank", f"Context words (top {CONTEXT_K})"],
        [
            [r["strategy"], r["method"], r["id"], r["type"], r["rank"],
             f"=IF(AND(E{i}>=1,E{i}<=1),1,0)", f"=IF(AND(E{i}>=1,E{i}<=3),1,0)",
             f"=IF(AND(E{i}>=1,E{i}<=5),1,0)", f"=IF(E{i}>0,1/E{i},0)", r["context_words"]]
            for i, r in enumerate(rows, start=2)
        ],
    )  # fmt: skip
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = f"A1:J{last}"
    set_widths(detail, [28, 26, 10, 12, 14, 8, 8, 8, 12, 14])

    def avg(column: str, *criteria: tuple[str, str]) -> str:
        conditions = ",".join(f"'Per question'!${c}$2:${c}${last},{v}" for c, v in criteria)
        return f"=IFERROR(AVERAGEIFS('Per question'!${column}$2:${column}${last},{conditions}),0)"

    # Summary: one row per strategy x retrieval method, every metric a formula.
    summary = workbook.create_sheet("Summary", 0)
    title(
        summary,
        "Retrieval evaluation: chunking strategies x retrieval methods",
        f"{len(questions)} questions. Recall@k = share of questions whose evidence is in the top k "
        f"results. Embedding model: {embedder_name}.",
    )
    headers = ["Strategy", "Retrieval", "Chunks", "Avg words per chunk", "Recall@1", "Recall@3",
               "Recall@5", "MRR", "Keyword questions Recall@5", "Paraphrase questions Recall@5",
               f"Avg context words (top {CONTEXT_K})"]  # fmt: skip

    def lookup(column: str, row: int) -> str:
        rows_ = f"$2:${column}${len(strategies) + 1}"
        names = f"Strategies!$A$2:$A${len(strategies) + 1}"
        return f"=INDEX(Strategies!${column}{rows_},MATCH($A{row},{names},0))"

    summary_rows = []
    for i, (strategy, method) in enumerate(((s, m) for s in strategies for m in methods), start=5):
        summary_rows.append([
            strategy.name, method, lookup("E", i), lookup("F", i),
            avg("F", ("A", f"$A{i}"), ("B", f"$B{i}")),
            avg("G", ("A", f"$A{i}"), ("B", f"$B{i}")),
            avg("H", ("A", f"$A{i}"), ("B", f"$B{i}")),
            avg("I", ("A", f"$A{i}"), ("B", f"$B{i}")),
            avg("H", ("A", f"$A{i}"), ("B", f"$B{i}"), ("D", '"keyword"')),
            avg("H", ("A", f"$A{i}"), ("B", f"$B{i}"), ("D", '"paraphrase"')),
            avg("J", ("A", f"$A{i}"), ("B", f"$B{i}")),
        ])  # fmt: skip
    write_table(summary, 4, headers, summary_rows)
    end = 4 + len(summary_rows)
    for row in summary.iter_rows(min_row=5, max_row=end, min_col=5, max_col=10):
        for cell in row:
            cell.number_format = "0.00" if cell.column == 8 else "0%"
    for row in summary.iter_rows(min_row=5, max_row=end, min_col=11, max_col=11):
        row[0].number_format = "#,##0"
    for row in summary.iter_rows(min_row=5, max_row=end, min_col=4, max_col=4):
        row[0].number_format = "0"
    summary.conditional_formatting.add(f"E5:J{end}", GREEN_SCALE)
    summary.freeze_panes = "C5"
    summary.auto_filter.ref = f"A4:K{end}"
    set_widths(summary, [28, 26, 9, 11, 10, 10, 10, 8, 14, 14, 14])
    summary.row_dimensions[4].height = 45

    # Matrix: strategies down, retrieval methods across, for Recall@5 and MRR.
    matrix = workbook.create_sheet("Matrix", 1)
    title(matrix, "Recall@5 and MRR by strategy and retrieval method", "Greener is better.")
    for block, (label, column, fmt) in enumerate((("Recall@5", "H", "0%"), ("MRR", "I", "0.00"))):
        top = 4 + block * (len(strategies) + 4)
        matrix.cell(row=top - 1, column=1, value=label).font = Font(name=FONT, bold=True)
        rows_ = []
        for r, strategy in enumerate(strategies, start=top + 1):
            rows_.append([strategy.name] + [
                avg(column, ("A", f"$A{r}"), ("B", f"{get_column_letter(c + 2)}${top}"))
                for c in range(len(methods))
            ])  # fmt: skip
        write_table(matrix, top, ["Strategy", *methods], rows_)
        last_col = get_column_letter(len(methods) + 1)
        span = f"B{top + 1}:{last_col}{top + len(strategies)}"
        for row in matrix[span]:
            for cell in row:
                cell.number_format = fmt
        matrix.conditional_formatting.add(span, GREEN_SCALE)
        matrix.row_dimensions[top].height = 45
    set_widths(matrix, [28] + [14] * len(methods))

    # Strategies, questions and settings: everything needed to reproduce the run.
    strategy_sheet = workbook.create_sheet("Strategies")
    write_table(
        strategy_sheet,
        1,
        ["Strategy", "Description", "Chunk size (words)", "Overlap (words)", "Chunks",
         "Avg words per chunk", "Heading path in searched text"],
        [[s.name, s.description, s.size, s.overlap, len(s.search_texts),
          round(mean(s.chunk_words.values()), 1),
          "yes" if s.headings_in_search_text else "no"] for s in strategies],
    )  # fmt: skip
    set_widths(strategy_sheet, [28, 70, 12, 12, 10, 12, 14])

    question_sheet = workbook.create_sheet("Questions")
    write_table(
        question_sheet,
        1,
        ["ID", "Type", "Question", "Evidence (any of)", "Where the answer is"],
        [[q["id"], q["type"], q["question"], " | ".join(q["evidence"]), q["section"]]
         for q in questions],
    )  # fmt: skip
    set_widths(question_sheet, [6, 12, 60, 90, 50])

    settings = workbook.create_sheet("Settings")
    write_table(
        settings,
        1,
        ["Setting", "Value"],
        [
            ["Run at (UTC)", datetime.now(UTC).strftime("%Y-%m-%d %H:%M")],
            ["Document", f"Serverless Applications Lens (July 14, 2022 edition), {SOURCE_URL}"],
            ["Embedding model", embedder_name],
            ["Results taken from each search before fusion", CANDIDATES],
            ["A question is missed if not answered within the top", MAX_RANK],
            ["Context sent to the model (top results)", CONTEXT_K],
            ["Weighted hybrid", "min-max scores per search, alpha x vector + (1 - alpha) x BM25"],
            ["Weighted hybrid alphas", ", ".join(str(a) for a in ALPHAS)],
            ["RRF", f"sum of 1 / ({RRF_K} + rank), equal weights"],
            ["BM25", "k1 = 1.5, b = 0.75, lowercase words, English stopwords removed"],
            ["Hit rule", "a retrieved text contains one of the question's evidence quotes "
             "(case and curly quotes ignored)"],
            ["Parent-child", "ranks count distinct parent sections, so its contexts are larger"],
            ["Caveat", "30 questions: differences of one or two questions (3-7%) are noise"],
        ],
    )  # fmt: skip
    set_widths(settings, [48, 110])

    workbook.calculation.fullCalcOnLoad = True  # Excel computes the formulas when opened
    return workbook


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--embedder", choices=("local", "titan", "none"), default="local")
    parser.add_argument(
        "--region", default="eu-central-1", help="Bedrock region for --embedder titan"
    )
    parser.add_argument("--source", type=Path, default=SOURCE_PDF)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    if not args.source.exists():
        raise SystemExit(f"Source PDF not found at {args.source}. See README: Local development.")
    questions = json.loads(QUESTIONS.read_text())["questions"]
    pages, headings = read_pdf(args.source.read_bytes())
    strategies = build_strategies(extract_sections(pages, headings))
    embedder = {"local": LocalEmbedder, "titan": lambda: TitanEmbedder(args.region),
                "none": lambda: None}[args.embedder]()  # fmt: skip

    rows = []
    for strategy in strategies:
        print(f"Evaluating {strategy.name} ({len(strategy.search_texts)} chunks)...")
        rows.extend(evaluate(strategy, questions, embedder))
    methods = list(dict.fromkeys(row["method"] for row in rows))

    output = args.output or RESULTS_DIR / f"retrieval_eval_{args.embedder}.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    name = embedder.name if embedder else "none (keyword search only)"
    build_workbook(strategies, methods, questions, rows, name).save(output)

    print(f"\n{'Strategy':30} {'Retrieval':26} {'R@1':>5} {'R@5':>5} {'MRR':>5}")
    for strategy in strategies:
        for method in methods:
            ranks = [
                r["rank"] for r in rows if r["strategy"] == strategy.name and r["method"] == method
            ]
            r1 = mean(1 <= r <= 1 for r in ranks)
            r5 = mean(1 <= r <= 5 for r in ranks)
            mrr = mean(1 / r if r else 0 for r in ranks)
            print(f"{strategy.name:30} {method:26} {r1:5.0%} {r5:5.0%} {mrr:5.2f}")
    print(f"\nWrote {output}")


if __name__ == "__main__":
    main()
